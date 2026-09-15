#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
NSAマッピング_マスター.xlsx から Web アプリ用の tech_items.json を生成するスクリプト。

読み込むシート:
  - 導入検討ツール   … No. / a.技術要素 / b.投資要素 / c.政策要素 / d.その他
  - 技術解説一覧     … No. / 分野・品目 / 食料安全保障要素 / 技術要素名 /
                        技術概要 / 期待される効果 / 導入にあたっての課題 / 出典・参考リンク

出力: tech_items.json
  {
    "generatedAt": "2026-09-14T12:00:00+09:00",
    "sourceFile": "NSAマッピング_マスター.xlsx",
    "itemCount": 199,
    "items": [
      {
        "id": 1,
        "cat": "分野横断：制度・基盤",
        "el": "A",
        "t": "農業投入財（種子・肥料）の流通・供給網整備、アグロディーラー支援",
        "tl": "中", "td": "物流・在庫管理、種子品質検査等の技術支援",
        "il": "大", "idt": "流通網・倉庫・アグロディーラー拠点整備",
        "pl": "必須", "pd": "流通事業者の参入規制緩和、種子・肥料の品質認証制度",
        "n": "道路インフラの整備状況に依存",
        "overview": "…",
        "effects": "…",
        "sources": [{"label": "FAO: Toolkit ...", "url": "https://..."}]
      }, ...
    ]
  }

使い方:
  python3 excel_to_json.py NSAマッピング_マスター.xlsx tech_items.json
"""
import json
import re
import sys
from datetime import datetime, timezone, timedelta

import openpyxl

FS_PREFIX_TO_CODE = {
    "Availability": "A",
    "Accessibility": "B",
    "Stability": "C",
    "Utilization": "D",
}

URL_RE = re.compile(r"(https?://\S+)$")
SEP_CHARS = " :：-–—　"


def split_level_detail(raw):
    """'中：普及員の技術研修…' -> ('中', '普及員の技術研修…')"""
    if not raw:
        return "", ""
    raw = str(raw).strip()
    for sep in ("：", ":"):
        if sep in raw:
            level, detail = raw.split(sep, 1)
            return level.strip(), detail.strip()
    return "", raw


def el_code_from_text(text):
    if not text:
        return ""
    for prefix, code in FS_PREFIX_TO_CODE.items():
        if text.startswith(prefix):
            return code
    return ""


def parse_sources(raw):
    """複数行の '出典: URL' 形式テキストを [{label, url}, ...] に分解する。"""
    if not raw:
        return []
    sources = []
    for line in str(raw).splitlines():
        line = line.strip()
        if not line:
            continue
        m = URL_RE.search(line)
        if m:
            url = m.group(1)
            label = line[: m.start()].rstrip(SEP_CHARS).strip()
            sources.append({"label": label or url, "url": url})
        else:
            sources.append({"label": line, "url": ""})
    return sources


def load_intro_sheet(ws):
    """導入検討ツール: No. -> {t, tl, td, il, idt, pl, pd, n}"""
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    out = {}
    for row in rows:
        no = row[0]
        if no is None:
            continue
        t = row[3]
        tl, td = split_level_detail(row[4])
        il, idt = split_level_detail(row[5])
        pl, pd = split_level_detail(row[6])
        n = (row[7] or "").strip() if row[7] else ""
        out[int(no)] = {
            "t": (t or "").strip(),
            "tl": tl, "td": td,
            "il": il, "idt": idt,
            "pl": pl, "pd": pd,
            "n": n,
        }
    return out


def load_explain_sheet(ws):
    """技術解説一覧: No. -> {cat, el, t, overview, effects, sources}"""
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    out = {}
    for row in rows:
        no = row[0]
        if no is None:
            continue
        cat = (row[1] or "").strip()
        el = el_code_from_text(row[2] or "")
        t = (row[3] or "").strip()
        overview = (row[4] or "").strip()
        effects = (row[5] or "").strip()
        sources = parse_sources(row[7])
        out[int(no)] = {
            "cat": cat, "el": el, "t": t,
            "overview": overview, "effects": effects,
            "sources": sources,
        }
    return out


def main():
    if len(sys.argv) < 3:
        print("usage: excel_to_json.py <master.xlsx> <out.json>", file=sys.stderr)
        sys.exit(1)
    src_path, out_path = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(src_path, data_only=True)
    intro = load_intro_sheet(wb["導入検討ツール"])
    explain = load_explain_sheet(wb["技術解説一覧"])

    ids = sorted(set(intro.keys()) | set(explain.keys()))
    mismatches = []
    items = []
    for i in ids:
        a = intro.get(i)
        b = explain.get(i)
        if a is None or b is None:
            mismatches.append((i, "missing in " + ("導入検討ツール" if a is None else "技術解説一覧")))
            continue
        if a["t"] and b["t"] and a["t"] != b["t"]:
            mismatches.append((i, f"title mismatch: {a['t'][:20]!r} vs {b['t'][:20]!r}"))
        items.append({
            "id": i,
            "cat": b["cat"],
            "el": b["el"],
            "t": a["t"] or b["t"],
            "tl": a["tl"], "td": a["td"],
            "il": a["il"], "idt": a["idt"],
            "pl": a["pl"], "pd": a["pd"],
            "n": a["n"],
            "overview": b["overview"],
            "effects": b["effects"],
            "sources": b["sources"],
        })

    JST = timezone(timedelta(hours=9))
    payload = {
        "generatedAt": datetime.now(JST).isoformat(timespec="seconds"),
        "sourceFile": src_path.split("/")[-1],
        "itemCount": len(items),
        "items": items,
    }

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    print(f"wrote {out_path}: {len(items)} items")
    if mismatches:
        print(f"WARNING: {len(mismatches)} inconsistencies found:", file=sys.stderr)
        for no, msg in mismatches[:20]:
            print(f"  No.{no}: {msg}", file=sys.stderr)


if __name__ == "__main__":
    main()
